#!/usr/bin/env python3
"""
md_to_docx.py — Convert manuscript markdown to a DOCX.
Uses python-docx. Chapter headings as Heading 2, body as serif paragraphs.

Usage:
    python3 md_to_docx.py <input.md> <output.docx>
"""
import sys
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


def parse_markdown(md: str):
    lines = md.split('\n')
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith('---'):
            yield ('hr', '')
            i += 1
            continue
        if line.startswith('# '):
            yield ('h1', line[2:].strip())
            i += 1
            continue
        if line.startswith('## '):
            yield ('h2', line[3:].strip())
            i += 1
            continue
        buf = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(('#', '---')):
            buf.append(lines[i].rstrip())
            i += 1
        yield ('p', ' '.join(buf))


def add_runs_with_inline(paragraph, text: str):
    """Parse **bold** and *italic* and add runs."""
    # Split on bold first
    parts = re.split(r'(\*\*.+?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            # Italic split
            sub = re.split(r'(?<!\*)(\*(?!\*)(?:[^*]+?)\*(?!\*))', part)
            for s in sub:
                if s.startswith('*') and s.endswith('*') and len(s) > 2 and not s.startswith('**'):
                    run = paragraph.add_run(s[1:-1])
                    run.italic = True
                else:
                    paragraph.add_run(s)


def build_docx(md_path: str, docx_path: str):
    md = Path(md_path).read_text(encoding='utf-8')
    doc = Document()
    # Body style — serif
    style = doc.styles['Normal']
    style.font.name = 'Georgia'
    style.font.size = Pt(11)

    for typ, text in parse_markdown(md):
        if typ == 'h1':
            p = doc.add_paragraph()
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt(24)
            r.font.name = 'Georgia'
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif typ == 'h2':
            p = doc.add_heading(text, level=2)
            for r in p.runs:
                r.font.name = 'Georgia'
        elif typ == 'p':
            if text.startswith('*') and text.endswith('*') and len(text) > 2:
                clean = text.lstrip('*').rstrip('*').strip()
                p = doc.add_paragraph()
                run = p.add_run(clean)
                run.italic = True
            else:
                p = doc.add_paragraph()
                add_runs_with_inline(p, text)
                p.paragraph_format.first_line_indent = Inches(0.25)
        elif typ == 'hr':
            doc.add_paragraph('_' * 40)

    doc.save(docx_path)


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print('Usage: md_to_docx.py <input.md> <output.docx>', file=sys.stderr)
        sys.exit(2)
    build_docx(sys.argv[1], sys.argv[2])
    print(f'[md_to_docx] wrote {sys.argv[2]}', file=sys.stderr)
